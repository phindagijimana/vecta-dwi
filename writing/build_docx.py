"""Convert writing/*.md sections to .docx files using python-docx.

Usage:
    python3 writing/build_docx.py

Outputs one .docx per section plus a combined manuscript.docx.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WRITING_DIR = Path(__file__).parent

SECTIONS = [
    ("abstract.md",      "Abstract"),
    ("introduction.md",  "Introduction"),
    ("methods.md",       "Methods"),
    ("results.md",       "Results"),
    ("discussion.md",    "Discussion"),
    ("references.md",    "References"),
]


def _set_normal_style(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    para_fmt = style.paragraph_format
    para_fmt.space_after = Pt(6)
    para_fmt.line_spacing = Pt(24)  # double-spaced


def _add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0] if h.runs else h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13) if level == 1 else Pt(12)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)


def _add_paragraph(doc, text):
    if not text.strip():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Handle inline code/bold markers simply: strip backticks, keep text
    clean = re.sub(r"`([^`]+)`", r"\1", text)
    clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", clean)
    run = p.add_run(clean)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def _add_table(doc, header_row, data_rows):
    n_cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=n_cols)
    table.style = "Table Grid"
    # Header
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header_row):
        hdr_cells[i].text = h.strip()
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.name = "Times New Roman"
    # Data rows
    for ri, row in enumerate(data_rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val.strip().lstrip("*").rstrip("*")
            for para in cells[ci].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Times New Roman"
    doc.add_paragraph()


def _parse_md_block(doc, text):
    """Parse a block of markdown text and add it to the doc."""
    lines = text.splitlines()
    i = 0
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("### "):
            if in_table:
                _flush_table(doc, table_lines)
                table_lines, in_table = [], False
            _add_heading(doc, line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith("## "):
            if in_table:
                _flush_table(doc, table_lines)
                table_lines, in_table = [], False
            _add_heading(doc, line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("# "):
            if in_table:
                _flush_table(doc, table_lines)
                table_lines, in_table = [], False
            _add_heading(doc, line[2:].strip(), level=1)
            i += 1
            continue

        # Tables
        if line.startswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                _flush_table(doc, table_lines)
                table_lines, in_table = [], False

        # Code blocks — skip fences, treat content as plain text
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                _add_paragraph(doc, lines[i])
                i += 1
            i += 1  # skip closing fence
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Continuation paragraphs — collect wrapped lines
        para_lines = []
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].startswith("|") and not lines[i].startswith("```"):
            para_lines.append(lines[i])
            i += 1
        if para_lines:
            _add_paragraph(doc, " ".join(para_lines))
        continue

    if in_table:
        _flush_table(doc, table_lines)


def _flush_table(doc, table_lines):
    """Parse collected pipe-table lines into a docx table."""
    rows = [l for l in table_lines if not re.match(r"^\|[-| ]+\|$", l)]
    if not rows:
        return
    parsed = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    if not parsed:
        return
    _add_table(doc, parsed[0], parsed[1:])


def build_section_docx(md_path: Path, title: str) -> Document:
    doc = Document()
    _set_normal_style(doc)

    # Title
    t = doc.add_heading(title, level=1)
    for run in t.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)

    text = md_path.read_text()
    # Strip top-level heading already used as title
    text = re.sub(r"^#\s+.+\n", "", text, count=1)

    _parse_md_block(doc, text)
    return doc


def build_combined_docx(sections) -> Document:
    doc = Document()
    _set_normal_style(doc)

    # Title page
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(
        "Vecta-DWI: A Declarative Framework for Data Birth Integrity\n"
        "Assessment of Diffusion MRI Datasets"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    doc.add_paragraph()

    for md_file, section_title in sections:
        md_path = WRITING_DIR / md_file
        if not md_path.is_file():
            print(f"  Skipping {md_file} (not found)")
            continue

        doc.add_page_break()
        h = doc.add_heading(section_title, level=1)
        for run in h.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)

        text = md_path.read_text()
        text = re.sub(r"^#\s+.+\n", "", text, count=1)
        _parse_md_block(doc, text)

    return doc


if __name__ == "__main__":
    print("Building section .docx files...")
    for md_file, title in SECTIONS:
        md_path = WRITING_DIR / md_file
        if not md_path.is_file():
            print(f"  Skipping {md_file} (not found)")
            continue
        doc = build_section_docx(md_path, title)
        out = WRITING_DIR / md_file.replace(".md", ".docx")
        doc.save(out)
        print(f"  Wrote {out.name}")

    print("Building combined manuscript.docx...")
    combined = build_combined_docx(SECTIONS)
    out = WRITING_DIR / "manuscript.docx"
    combined.save(out)
    print(f"  Wrote {out.name}")
    print("Done.")
